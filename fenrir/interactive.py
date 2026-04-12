from __future__ import annotations

import copy
import importlib
import json
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.base import clone
from sklearn.preprocessing import MinMaxScaler, RobustScaler, StandardScaler

from .core import Fenrir

try:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    Font = None
    PatternFill = None
    get_column_letter = None

try:
    import ipywidgets as widgets
    from IPython import get_ipython
    from IPython.display import clear_output, display
except Exception:
    widgets = None
    clear_output = None
    display = None
    get_ipython = None


def _publish_notebook_bindings(**values):
    if get_ipython is None:
        return

    shell = get_ipython()
    if shell is None or not hasattr(shell, "user_ns"):
        return

    shell.user_ns.update(values)


def _require_widgets():
    if widgets is None or clear_output is None or display is None:
        raise ImportError(
            "ipywidgets e IPython son necesarios para la capa interactiva de Fenrir. "
            "Instala el extra notebook con 'pip install -e .[notebook]'."
        )


def _require_matplotlib_pyplot():
    try:
        return importlib.import_module("matplotlib.pyplot")
    except ImportError as exc:
        raise ImportError(
            "matplotlib es necesario para las funciones graficas de Fenrir. "
            "Instalalo con 'pip install matplotlib'."
        ) from exc


def _round_frame(frame, round_digits):
    if round_digits is None:
        return frame

    rounded = frame.copy()
    numeric_columns = rounded.select_dtypes(include=[np.number]).columns
    rounded.loc[:, numeric_columns] = rounded.loc[:, numeric_columns].round(int(round_digits))
    return rounded


def _key_metric_view(model):
    metrics = model.metric_report().copy()
    selected = metrics[
        metrics["metric"].isin(
            [
                "explained_variance",
                "variance_first_two",
                "silhouette",
                "davies_bouldin",
                "ari",
                "nmi",
                "v_measure",
            ]
        )
    ]
    return selected.loc[
        :,
        [
            column
            for column in ["metric", "metric_label", "value", "value_display", "interpretation"]
            if column in selected.columns
        ],
    ]


def _default_scaler_catalog():
    return {
        "StandardScaler": StandardScaler(),
        "RobustScaler": RobustScaler(),
        "MinMaxScaler": MinMaxScaler(),
    }


def _normalize_export_path(file_name, default_stem, file_format):
    normalized_format = "txt" if file_format == "txt-list" else file_format
    raw_name = str(file_name or "").strip()
    path = Path(raw_name) if raw_name else Path.home() / "Downloads" / f"{default_stem}.{normalized_format}"
    if not path.is_absolute():
        path = Path.home() / "Downloads" / path

    expected_suffix = f".{normalized_format}"
    if path.suffix.lower() != expected_suffix:
        path = path.with_suffix(expected_suffix)

    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _export_payload(frame=None, list_values=None, file_name="", default_stem="fenrir_export", file_format="xlsx"):
    path = _normalize_export_path(file_name, default_stem, file_format)
    export_frame = None if frame is None else frame.copy()

    if file_format == "xlsx":
        if export_frame is None:
            export_frame = pd.DataFrame({"value": list(list_values or [])})
        export_frame.to_excel(path, index=True)
        return path

    if file_format == "csv":
        if export_frame is None:
            export_frame = pd.DataFrame({"value": list(list_values or [])})
        export_frame.to_csv(path, index=True, encoding="utf-8-sig")
        return path

    if file_format == "json":
        if list_values is not None and export_frame is None:
            path.write_text(
                pd.Series(list(list_values)).to_json(orient="values", force_ascii=False, indent=2),
                encoding="utf-8",
            )
        else:
            export_frame = pd.DataFrame() if export_frame is None else export_frame
            path.write_text(export_frame.to_json(orient="records", force_ascii=False, indent=2), encoding="utf-8")
        return path

    values = list(list_values) if list_values is not None else None
    if values is None:
        if export_frame is None:
            values = []
        elif export_frame.shape[1] == 1:
            values = export_frame.iloc[:, 0].astype(str).tolist()
        else:
            values = export_frame.index.astype(str).tolist()

    path.write_text("\n".join(str(value) for value in values), encoding="utf-8")
    return path


def _build_export_controls(default_name, button_label):
    file_widget = widgets.Text(
        value=default_name,
        description="Archivo",
        layout=widgets.Layout(width="310px"),
    )
    format_widget = widgets.Dropdown(
        options=[
            ("Excel", "xlsx"),
            ("CSV", "csv"),
            ("JSON", "json"),
            ("TXT lista", "txt-list"),
        ],
        value="xlsx",
        description="Formato",
        layout=widgets.Layout(width="220px"),
    )
    button = widgets.Button(description=button_label, icon="download")
    return file_widget, format_widget, button


def _export_figure(fig, file_name="", default_stem="fenrir_plot", file_format="png"):
    path = _normalize_export_path(file_name, default_stem, file_format)
    fig.savefig(path, dpi=180, bbox_inches="tight")
    return path


def _render_plot_figure(model, plot_name, scaler_name=None, top_n=10, width=8, height=5):
    plt_module = _require_matplotlib_pyplot()
    scaler_name = scaler_name or model.best_configuration()["scaler_name"]
    figsize = (float(width), float(height))
    if plot_name == "correlation_circle":
        figsize = (max(figsize[0], 7.0), max(figsize[1], 7.0))

    fig, ax = plt_module.subplots(figsize=figsize)

    if plot_name == "scree":
        model.plot_scree(scaler_name=scaler_name, ax=ax)
    elif plot_name == "correlation_circle":
        model.plot_correlation_circle(scaler_name=scaler_name, top_n=int(top_n), ax=ax)
    elif plot_name == "feature_contributions":
        model.plot_feature_contributions(scaler_name=scaler_name, top_n=int(top_n), ax=ax)
    elif plot_name == "feature_quality":
        model.plot_feature_quality(scaler_name=scaler_name, top_n=int(top_n), ax=ax)
    elif plot_name == "clusters":
        model.plot_clusters(scaler_name=scaler_name, ax=ax)
    else:
        plt_module.close(fig)
        raise ValueError(f"Grafico no soportado: {plot_name}")

    fig.tight_layout()
    return fig


def _workbench_presets_for(model):
    presets = getattr(model, "workbench_presets_", None)
    if presets is None:
        presets = {}
        model.workbench_presets_ = presets
    return presets


def _save_workbench_preset(self, name, fenrir_config, analysis_config):
    preset_name = str(name or "").strip()
    if not preset_name:
        raise ValueError("El nombre del preset no puede estar vacio.")

    payload = {
        "fenrir_config": copy.deepcopy(dict(fenrir_config)),
        "analysis_config": copy.deepcopy(dict(analysis_config)),
    }
    _workbench_presets_for(self)[preset_name] = payload
    return copy.deepcopy(payload)


def _get_workbench_preset(self, name):
    preset = _workbench_presets_for(self).get(name)
    return None if preset is None else copy.deepcopy(preset)


def _delete_workbench_preset(self, name):
    removed = _workbench_presets_for(self).pop(name, None)
    return None if removed is None else copy.deepcopy(removed)


def _list_workbench_presets(self):
    return sorted(_workbench_presets_for(self))


def _json_safe_value(value):
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, range):
        return [int(item) for item in value]
    if isinstance(value, (list, tuple, set)):
        return [_json_safe_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _json_safe_value(item) for key, item in value.items()}
    return str(value)


def _normalize_preset_payload(payload):
    fenrir_payload = copy.deepcopy(dict(payload.get("fenrir_config", {})))
    analysis_payload = copy.deepcopy(dict(payload.get("analysis_config", {})))

    scalers = fenrir_payload.get("scalers")
    if isinstance(scalers, dict):
        fenrir_payload["scalers"] = list(scalers.keys())
    elif scalers is not None:
        fenrir_payload["scalers"] = list(scalers)

    cluster_range = fenrir_payload.get("cluster_range")
    if cluster_range is not None:
        fenrir_payload["cluster_range"] = [int(value) for value in cluster_range]

    for key in ["features", "exclude_columns", "algorithms"]:
        if fenrir_payload.get(key) is not None:
            fenrir_payload[key] = list(fenrir_payload[key])

    return {
        "fenrir_config": _json_safe_value(fenrir_payload),
        "analysis_config": _json_safe_value(analysis_payload),
    }


def _restore_preset_payload(payload):
    fenrir_payload = copy.deepcopy(dict(payload.get("fenrir_config", {})))
    analysis_payload = copy.deepcopy(dict(payload.get("analysis_config", {})))

    for key in ["features", "exclude_columns", "algorithms", "scalers"]:
        if fenrir_payload.get(key) is not None:
            fenrir_payload[key] = tuple(fenrir_payload[key])

    cluster_range = fenrir_payload.get("cluster_range")
    if cluster_range is not None:
        fenrir_payload["cluster_range"] = tuple(int(value) for value in cluster_range)

    return {
        "fenrir_config": fenrir_payload,
        "analysis_config": analysis_payload,
    }


def _preset_store_to_json_payload(preset_store):
    return {
        str(name): _normalize_preset_payload(payload)
        for name, payload in dict(preset_store).items()
    }


def _preset_store_from_json_payload(payload):
    return {
        str(name): _restore_preset_payload(item)
        for name, item in dict(payload).items()
    }


def _save_preset_store_json(preset_store, file_name="", default_stem="fenrir_workbench_presets"):
    path = _normalize_export_path(file_name, default_stem, "json")
    json_payload = {
        "version": 1,
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "presets": _preset_store_to_json_payload(preset_store),
    }
    path.write_text(json.dumps(json_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _load_preset_store_json(file_name="", default_stem="fenrir_workbench_presets"):
    path = _normalize_export_path(file_name, default_stem, "json")
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de presets: {path}")

    raw_payload = json.loads(path.read_text(encoding="utf-8"))
    preset_payload = raw_payload.get("presets", raw_payload)
    return _preset_store_from_json_payload(preset_payload), path


def _save_workbench_presets_json(self, file_name="", default_stem="fenrir_workbench_presets"):
    return _save_preset_store_json(_workbench_presets_for(self), file_name=file_name, default_stem=default_stem)


def _load_workbench_presets_json(self, file_name="", default_stem="fenrir_workbench_presets", merge=False):
    loaded_store, path = _load_preset_store_json(file_name=file_name, default_stem=default_stem)
    if merge:
        merged_store = copy.deepcopy(_workbench_presets_for(self))
        merged_store.update(loaded_store)
        loaded_store = merged_store

    self.workbench_presets_ = copy.deepcopy(loaded_store)
    return copy.deepcopy(loaded_store), path


def _excel_safe_value(value):
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(_json_safe_value(value), ensure_ascii=False)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        return value.item()
    return value


def _prepare_excel_frame(frame):
    export_frame = pd.DataFrame() if frame is None else frame.copy()
    if isinstance(export_frame.columns, pd.MultiIndex):
        export_frame.columns = [
            " | ".join(str(part) for part in column_value if str(part) not in {"", "None"})
            for column_value in export_frame.columns
        ]
    else:
        export_frame.columns = [_excel_safe_value(column) for column in export_frame.columns]
    if isinstance(export_frame.index, pd.MultiIndex):
        export_frame.index = [" | ".join(str(part) for part in index_value) for index_value in export_frame.index]
    else:
        export_frame.index = [_excel_safe_value(index_value) for index_value in export_frame.index]

    for column in export_frame.columns:
        export_frame[column] = export_frame[column].map(_excel_safe_value)
    return export_frame


def _config_sheet(fenrir_config, analysis_config):
    rows = []
    for key, value in dict(fenrir_config or {}).items():
        rows.append({"bloque": "fenrir", "parametro": key, "valor": _excel_safe_value(value)})
    for key, value in dict(analysis_config or {}).items():
        rows.append({"bloque": "analysis", "parametro": key, "valor": _excel_safe_value(value)})
    return pd.DataFrame(rows)


def _executive_summary_sheet(model, report=None):
    rows = []
    best_config = model.best_configuration()
    key_map = {
        "scaler_name": "scaler_base",
        "algorithm": "algoritmo_base",
        "n_clusters": "clusters_base",
        "n_components": "componentes_base",
        "explained_variance": "varianza_pca",
        "variance_first_two": "varianza_pc1_pc2",
        "silhouette": "silhouette",
        "davies_bouldin": "davies_bouldin",
    }
    for key, value in best_config.items():
        rows.append(
            {
                "bloque": "best_configuration",
                "campo": key_map.get(key, key),
                "valor": _excel_safe_value(value),
            }
        )

    metric_frame = _key_metric_view(model)
    for row in metric_frame.itertuples(index=False):
        rows.append(
            {
                "bloque": "metric_report",
                "campo": getattr(row, "metric_label", row.metric),
                "valor": _excel_safe_value(getattr(row, "value_display", row.value)),
                "detalle": getattr(row, "interpretation", ""),
            }
        )

    if report is not None:
        rows.append(
            {
                "bloque": "compactos",
                "campo": "best_compact_variables",
                "valor": ", ".join(str(item) for item in report.get("best_compact_variables", [])),
            }
        )
        holdout_top = report.get("holdout_top")
        if holdout_top is not None and not holdout_top.empty:
            top_row = holdout_top.iloc[0]
            rows.extend(
                [
                    {
                        "bloque": "validacion",
                        "campo": "mejor_holdout",
                        "valor": _excel_safe_value(top_row.get("configuration", "")),
                    },
                    {
                        "bloque": "validacion",
                        "campo": "lectura_holdout",
                        "valor": _excel_safe_value(top_row.get("holdout_silhouette_reading", "")),
                    },
                ]
            )
        stability_summary = report.get("stability_summary")
        if stability_summary is not None and not stability_summary.empty:
            stability_row = stability_summary.iloc[0]
            rows.append(
                {
                    "bloque": "estabilidad",
                    "campo": "lectura_estabilidad",
                    "valor": _excel_safe_value(stability_row.get("ari_vs_base_reading", stability_row.get("stability_reading", ""))),
                }
            )

    return pd.DataFrame(rows)


def _excel_column_widths(frame: pd.DataFrame) -> list[int]:
    prepared = pd.DataFrame() if frame is None else frame.copy()
    index_header = "indice" if prepared.index.name is None else str(prepared.index.name)
    index_values = [str(value) for value in prepared.index.tolist()]
    widths = [min(max([len(index_header)] + [len(str(value)) for value in index_values], default=12) + 2, 42)]

    for column in prepared.columns:
        values = prepared[column].tolist() if len(prepared) else []
        width = min(max([len(str(column))] + [len(str(value)) for value in values], default=12) + 2, 48)
        widths.append(width)
    return widths


def _format_excel_sheet(writer, sheet_name: str, frame: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    widths = _excel_column_widths(frame)
    max_row = len(frame.index)
    max_col = len(frame.columns)

    if hasattr(writer, "engine") and writer.engine == "xlsxwriter":
        workbook = writer.book
        header_format = workbook.add_format(
            {
                "bold": True,
                "font_color": "#FFFFFF",
                "bg_color": "#12324A",
                "border": 0,
            }
        )
        worksheet.freeze_panes(1, 1)
        worksheet.autofilter(0, 0, max_row, max_col)
        worksheet.set_row(0, None, header_format)
        for idx, width in enumerate(widths):
            worksheet.set_column(idx, idx, width)
        return

    if get_column_letter is None:
        return

    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions
    if Font is not None and PatternFill is not None:
        header_fill = PatternFill(fill_type="solid", fgColor="12324A")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
    for idx, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(idx)].width = width


def _ordered_executive_sheets(model, report, fenrir_config, current_analysis, best_scaler, top_n_value):
    sheets = [
        ("Resumen ejecutivo", _executive_summary_sheet(model, report=report)),
        ("Configuracion", _config_sheet(fenrir_config, current_analysis)),
        ("Leaderboard", model.summary(top_n=int(current_analysis["leaderboard_top_n"]))),
        ("Metricas clave", model.metric_report()),
        ("Tamano clusters", model.cluster_size_report()),
        ("Perfil clusters", model.cluster_groupby(top_n=top_n_value)),
        ("PCA componentes", model.pca_report(scaler_name=best_scaler)),
        ("PCA variables", model.source_feature_report(scaler_name=best_scaler, top_n=max(10, top_n_value))),
        ("Holdout resumen", report["holdout_summary"]),
        ("Holdout top", report["holdout_top"]),
        ("Compactos ranking", report["compact_subset_search"]),
        ("Compacto groupby", report.get("best_compact_groupby")),
        ("Preprocesado", report["preprocessing_report"]),
        ("Estabilidad resumen", report.get("stability_summary")),
        ("Estabilidad top", report.get("stability_top")),
    ]

    if getattr(model, "target_", None) is not None:
        sheets.insert(6, ("Cluster vs target", model.cluster_target_report(normalize="index")))

    return [(name, frame) for name, frame in sheets if frame is not None]


def _export_executive_workbook(self, file_name="", analysis_config=None, fenrir_config=None, default_stem="fenrir_executivo"):
    self._require_fit()
    current_analysis = _default_analysis_config(self, analysis_config)
    report = self.analysis_report(**current_analysis)
    best_scaler = self.best_configuration()["scaler_name"]
    top_n_value = int(current_analysis["top_n"])
    workbook_path = _normalize_export_path(file_name, default_stem, "xlsx")

    sheets = _ordered_executive_sheets(self, report, fenrir_config, current_analysis, best_scaler, top_n_value)

    with pd.ExcelWriter(workbook_path) as writer:
        for sheet_name, sheet_frame in sheets:
            prepared_frame = _prepare_excel_frame(sheet_frame)
            prepared_frame.to_excel(writer, sheet_name=sheet_name[:31], index=True)
            _format_excel_sheet(writer, sheet_name[:31], prepared_frame)

    return workbook_path


def _prepare_workbench_initial_configs(model, data=None, fenrir_config=None, analysis_config=None):
    raw_data = model._as_frame(model.data if data is None else data)
    available_columns = list(raw_data.columns)
    available_scalers = tuple(_default_scaler_catalog().keys())

    current_target = model.target if isinstance(model.target, str) and model.target in available_columns else None
    prepared_fenrir = {
        "target": current_target,
        "features": [] if model.features is None else [column for column in model.features if column in available_columns],
        "exclude_columns": [column for column in sorted(model.exclude_columns) if column in available_columns],
        "variance_threshold": float(model.variance_threshold),
        "cluster_range": tuple(model.cluster_range),
        "algorithms": tuple(model.algorithms),
        "scalers": tuple(name for name in model.scalers.keys() if name in available_scalers),
        "max_categories": int(model.max_categories),
        "random_state": int(model.random_state),
    }
    if fenrir_config:
        prepared_fenrir.update(copy.deepcopy(dict(fenrir_config)))

    prepared_analysis = copy.deepcopy(dict(analysis_config or {}))

    target_value = prepared_fenrir.get("target")
    prepared_fenrir["target"] = target_value if target_value in available_columns else None

    features_value = prepared_fenrir.get("features")
    if features_value is None:
        prepared_fenrir["features"] = []
    else:
        prepared_fenrir["features"] = [value for value in features_value if value in available_columns]

    exclude_value = prepared_fenrir.get("exclude_columns")
    if exclude_value is None:
        prepared_fenrir["exclude_columns"] = []
    else:
        prepared_fenrir["exclude_columns"] = [value for value in exclude_value if value in available_columns]

    cluster_value = prepared_fenrir.get("cluster_range")
    if cluster_value is None:
        prepared_fenrir["cluster_range"] = tuple(model.cluster_range)
    else:
        prepared_fenrir["cluster_range"] = tuple(int(value) for value in cluster_value)

    algorithms_value = prepared_fenrir.get("algorithms")
    if algorithms_value is not None:
        prepared_fenrir["algorithms"] = tuple(algorithms_value)

    scalers_value = prepared_fenrir.get("scalers")
    if isinstance(scalers_value, dict):
        normalized_scalers = tuple(name for name in scalers_value.keys() if name in available_scalers)
    elif scalers_value is None:
        normalized_scalers = tuple(name for name in model.scalers.keys() if name in available_scalers)
    else:
        normalized_scalers = tuple(name for name in scalers_value if name in available_scalers)
    prepared_fenrir["scalers"] = normalized_scalers or tuple(name for name in model.scalers.keys() if name in available_scalers)

    return prepared_fenrir, prepared_analysis


def _default_analysis_config(model, analysis_config=None):
    config = {
        "top_n": 10,
        "leaderboard_top_n": 10,
        "holdout_top_n": 10,
        "compact_max_features": min(5, max(2, len(model.list_influential_variables(top_n=None)))),
        "compact_criterion": "holdout_silhouette",
        "compact_include_holdout": None,
        "test_size": 0.25,
        "n_splits": 2,
        "stratify_target": True,
        "min_cluster_share": 0.01,
        "include_stability": True,
        "stability_sample_fraction": 0.8,
        "stability_n_splits": 3,
        "stability_top_n": 10,
        "stability_stratify_target": True,
    }
    if analysis_config:
        config.update(dict(analysis_config))

    config["top_n"] = int(config["top_n"])
    config["leaderboard_top_n"] = int(config.get("leaderboard_top_n", config["top_n"]))
    config["holdout_top_n"] = int(config.get("holdout_top_n", config["top_n"]))
    config["compact_max_features"] = int(config["compact_max_features"])
    config["test_size"] = float(config["test_size"])
    config["n_splits"] = int(config["n_splits"])
    config["min_cluster_share"] = float(config["min_cluster_share"])
    config["include_stability"] = bool(config["include_stability"])
    config["stability_sample_fraction"] = float(config["stability_sample_fraction"])
    config["stability_n_splits"] = int(config["stability_n_splits"])
    config["stability_top_n"] = int(config.get("stability_top_n", config["top_n"]))
    return config


def _build_model_from_config(base_model, fenrir_config, data=None):
    raw_data = base_model._as_frame(base_model.data if data is None else data)
    normalized_fenrir, _ = _prepare_workbench_initial_configs(
        base_model,
        data=data,
        fenrir_config=fenrir_config,
        analysis_config=None,
    )
    scaler_catalog = _default_scaler_catalog()
    selected_scalers = tuple(name for name in normalized_fenrir.get("scalers", ()) if name in scaler_catalog)
    if not selected_scalers:
        selected_scalers = tuple(name for name in base_model.scalers.keys() if name in scaler_catalog)

    cluster_values = tuple(normalized_fenrir.get("cluster_range") or tuple(base_model.cluster_range))

    model = Fenrir(
        raw_data.copy(),
        target=normalized_fenrir.get("target"),
        features=list(normalized_fenrir["features"]) or None,
        exclude_columns=list(normalized_fenrir.get("exclude_columns", [])),
        variance_threshold=float(normalized_fenrir.get("variance_threshold", base_model.variance_threshold)),
        cluster_range=range(int(min(cluster_values)), int(max(cluster_values)) + 1),
        algorithms=tuple(normalized_fenrir.get("algorithms") or base_model.algorithms),
        scalers={name: clone(scaler_catalog[name]) for name in selected_scalers},
        max_categories=int(normalized_fenrir.get("max_categories", base_model.max_categories)),
        random_state=int(normalized_fenrir.get("random_state", base_model.random_state)),
    )
    model.fit()
    return model, normalized_fenrir


def _flat_config_frame(fenrir_config, analysis_config, label):
    frame = _config_sheet(fenrir_config, analysis_config).copy()
    if frame.empty:
        return pd.DataFrame(columns=[label])

    index = frame["bloque"].astype(str) + "." + frame["parametro"].astype(str)
    values = frame["valor"].map(_excel_safe_value)
    return pd.DataFrame({label: values.values}, index=index)


def _validation_summary_frame(candidate):
    report = candidate["report"]
    holdout_top = report.get("holdout_top")
    holdout_row = holdout_top.iloc[0] if holdout_top is not None and not holdout_top.empty else pd.Series(dtype=object)
    stability_summary = report.get("stability_summary")
    stability_row = stability_summary.iloc[0] if stability_summary is not None and not stability_summary.empty else pd.Series(dtype=object)

    summary = {
        "best_compact_variables": ", ".join(str(item) for item in report.get("best_compact_variables", [])),
        "holdout_scaler": holdout_row.get("scaler_name"),
        "holdout_algorithm": holdout_row.get("algorithm"),
        "holdout_n_clusters": holdout_row.get("n_clusters"),
        "holdout_n_components": holdout_row.get("n_components"),
        "holdout_test_silhouette_mean": holdout_row.get("test_silhouette_mean"),
        "holdout_test_davies_bouldin_mean": holdout_row.get("test_davies_bouldin_mean"),
        "holdout_test_v_measure_mean": holdout_row.get("test_v_measure_mean"),
        "top_variables": ", ".join(candidate["model"].list_influential_variables(top_n=5)),
    }
    if not stability_row.empty:
        for column in stability_row.index:
            summary[f"stability.{column}"] = stability_row[column]

    return pd.DataFrame({candidate["label"]: pd.Series(summary)})


def _resolve_comparison_candidate(
    base_model,
    candidate_name,
    preset_store,
    current_fenrir_config=None,
    current_analysis_config=None,
    data=None,
):
    if candidate_name == "__current__":
        if current_fenrir_config is None:
            raise ValueError("Falta la configuracion actual para comparar el estado del workbench.")

        label = "Actual"
        model, normalized_fenrir = _build_model_from_config(base_model, current_fenrir_config, data=data)
        normalized_analysis = _default_analysis_config(model, current_analysis_config)
    else:
        preset = dict(preset_store or {}).get(candidate_name)
        if preset is None:
            raise ValueError(f"No existe el preset: {candidate_name}")

        label = str(candidate_name)
        model, normalized_fenrir = _build_model_from_config(base_model, preset.get("fenrir_config", {}), data=data)
        normalized_analysis = _default_analysis_config(model, preset.get("analysis_config", {}))

    report = model.analysis_report(**normalized_analysis)
    return {
        "label": label,
        "model": model,
        "fenrir_config": normalized_fenrir,
        "analysis_config": normalized_analysis,
        "report": report,
    }


def _compare_workbench_candidates(
    self,
    left_candidate,
    right_candidate,
    preset_store=None,
    current_fenrir_config=None,
    current_analysis_config=None,
    data=None,
):
    if left_candidate == right_candidate:
        raise ValueError("Selecciona dos configuraciones distintas para comparar.")

    store = copy.deepcopy(_workbench_presets_for(self) if preset_store is None else preset_store)
    left = _resolve_comparison_candidate(
        self,
        left_candidate,
        store,
        current_fenrir_config=current_fenrir_config,
        current_analysis_config=current_analysis_config,
        data=data,
    )
    right = _resolve_comparison_candidate(
        self,
        right_candidate,
        store,
        current_fenrir_config=current_fenrir_config,
        current_analysis_config=current_analysis_config,
        data=data,
    )

    configuration = pd.concat(
        [
            _flat_config_frame(left["fenrir_config"], left["analysis_config"], left["label"]),
            _flat_config_frame(right["fenrir_config"], right["analysis_config"], right["label"]),
        ],
        axis=1,
    )
    configuration.insert(
        0,
        "coincide",
        np.where(
            configuration[left["label"]].astype(str) == configuration[right["label"]].astype(str),
            "si",
            "no",
        ),
    )

    best_configuration = pd.concat(
        [
            left["model"].best_configuration().rename(left["label"]),
            right["model"].best_configuration().rename(right["label"]),
        ],
        axis=1,
    )

    quality = pd.concat(
        [
            _key_metric_view(left["model"]).set_index("metric")[["value"]].rename(columns={"value": left["label"]}),
            _key_metric_view(right["model"]).set_index("metric")[["value"]].rename(columns={"value": right["label"]}),
        ],
        axis=1,
    )

    validation = pd.concat(
        [
            _validation_summary_frame(left),
            _validation_summary_frame(right),
        ],
        axis=1,
    )

    return {
        "left_label": left["label"],
        "right_label": right["label"],
        "configuration": configuration,
        "best_configuration": best_configuration,
        "quality": quality,
        "validation": validation,
        "left_report": left["report"],
        "right_report": right["report"],
    }


def _pca_interactive_table(self, top_n=10, max_top_n=25):
    self._require_fit()
    _require_widgets()

    scaler_names = list(self.pca_profiles_.keys())
    max_components = max(len(profile["component_report"]) for profile in self.pca_profiles_.values())
    default_component_count = min(3, max(2, int(self.best_configuration()["n_components"])))

    scaler_widget = widgets.Dropdown(
        options=scaler_names,
        value=self.best_configuration()["scaler_name"],
        description="Scaler",
        layout=widgets.Layout(width="240px"),
    )
    view_widget = widgets.Dropdown(
        options=[
            ("Resumen de componentes", "component_report"),
            ("Variables originales", "source_feature_report"),
            ("Variables codificadas", "feature_report"),
            ("Lista influyente", "influential_list"),
            ("Scores PCA", "scores_preview"),
        ],
        value="source_feature_report",
        description="Vista",
        layout=widgets.Layout(width="260px"),
    )
    component_widget = widgets.SelectMultiple(
        options=[(f"PC{value}", value) for value in range(1, max_components + 1)],
        value=tuple(range(1, default_component_count + 1)),
        description="PCs",
        rows=min(8, max_components),
        layout=widgets.Layout(width="220px", height="180px"),
    )
    top_n_widget = widgets.IntSlider(
        value=min(max(3, int(top_n)), max_top_n),
        min=3,
        max=max_top_n,
        step=1,
        description="Top n",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    score_rows_widget = widgets.IntSlider(
        value=8,
        min=3,
        max=25,
        step=1,
        description="Filas",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    sort_widget = widgets.Dropdown(
        options=[
            ("Sin ordenar", None),
            ("Contribucion ponderada", "weighted_contribution"),
            ("Calidad ponderada", "weighted_cos2"),
            ("Contribucion total", "contribution_total"),
            ("Calidad total", "cos2_total"),
            ("Varianza explicada", "explained_variance"),
            ("Varianza acumulada", "cumulative_variance"),
        ],
        value="weighted_contribution",
        description="Orden",
        layout=widgets.Layout(width="280px"),
    )
    ascending_widget = widgets.Checkbox(value=False, description="Ascendente")
    round_widget = widgets.Dropdown(
        options=[("Sin redondeo", None)] + [(str(value), value) for value in range(0, 7)],
        value=4,
        description="Round",
        layout=widgets.Layout(width="220px"),
    )
    export_name_widget, export_format_widget, export_button = _build_export_controls(
        "fenrir_pca.xlsx",
        "Exportar PCA",
    )
    refresh_button = widgets.Button(
        description="Actualizar PCA",
        button_style="primary",
        icon="refresh",
    )

    status_output = widgets.Output()
    summary_output = widgets.Output()
    table_output = widgets.Output()
    export_output = widgets.Output()
    last_payload = {"frame": None, "list_values": None, "default_stem": "fenrir_pca"}

    def _sync_components(*_):
        profile = self._profile_for(scaler_name=scaler_widget.value)
        available = len(profile["component_report"])
        options = [(f"PC{value}", value) for value in range(1, available + 1)]
        previous = tuple(value for value in component_widget.value if value <= available)
        if not previous:
            default_limit = min(3, max(2, int(profile["selected_components"])))
            previous = tuple(range(1, default_limit + 1))
        component_widget.options = options
        component_widget.value = previous

    def _render(_=None):
        with status_output:
            clear_output(wait=True)
            print("Actualizando tabla PCA...")

        try:
            selected_components = list(component_widget.value) or None
            profile = self._profile_for(scaler_name=scaler_widget.value)
            component_report = profile["component_report"].copy()
            selected_total = (
                component_report.set_index("component").loc[selected_components, "explained_variance"].sum()
                if selected_components
                else component_report["explained_variance"].sum()
            )
            summary = pd.DataFrame(
                {
                    "valor": [
                        scaler_widget.value,
                        int(profile["selected_components"]),
                        list(component_widget.value),
                        float(selected_total),
                        float(profile["variance_first_two"]),
                    ]
                },
                index=[
                    "scaler_activo",
                    "pcs_sugeridas_por_varianza",
                    "pcs_consultadas",
                    "varianza_pcs_consultadas",
                    "varianza_pc1_pc2",
                ],
            )

            list_values = None
            if view_widget.value == "component_report":
                frame = self.pca_report(scaler_name=scaler_widget.value)
            elif view_widget.value == "source_feature_report":
                frame = self.source_feature_report(
                    scaler_name=scaler_widget.value,
                    components=selected_components,
                    top_n=int(top_n_widget.value),
                )
            elif view_widget.value == "feature_report":
                frame = self.feature_report(
                    scaler_name=scaler_widget.value,
                    components=selected_components,
                    top_n=int(top_n_widget.value),
                )
            elif view_widget.value == "influential_list":
                list_values = self.list_influential_variables(
                    scaler_name=scaler_widget.value,
                    components=selected_components,
                    top_n=int(top_n_widget.value),
                    grouped=True,
                )
                frame = pd.DataFrame({"rank": np.arange(1, len(list_values) + 1), "variable": list_values})
            else:
                n_components = max(selected_components) if selected_components else int(profile["selected_components"])
                frame = self.transform(scaler_name=scaler_widget.value, n_components=n_components).head(int(score_rows_widget.value))

            sort_column = sort_widget.value
            if sort_column is not None and sort_column in frame.columns:
                frame = frame.sort_values(sort_column, ascending=bool(ascending_widget.value))
            frame = _round_frame(frame, round_widget.value)

            last_payload["frame"] = frame.copy()
            last_payload["list_values"] = None if list_values is None else list(list_values)
            last_payload["default_stem"] = f"fenrir_pca_{view_widget.value}"

            with summary_output:
                clear_output(wait=True)
                display(summary)

            with table_output:
                clear_output(wait=True)
                display(frame)

            with status_output:
                clear_output(wait=True)
                print("Tabla PCA lista")
        except Exception as exc:
            last_payload["frame"] = None
            last_payload["list_values"] = None
            with summary_output:
                clear_output(wait=True)
            with table_output:
                clear_output(wait=True)
            with status_output:
                clear_output(wait=True)
                print(f"No se pudo generar la tabla PCA: {exc}")

    def _export(_=None):
        with export_output:
            clear_output(wait=True)
        try:
            if last_payload["frame"] is None:
                raise RuntimeError("Actualiza la tabla PCA antes de exportar.")
            path = _export_payload(
                frame=last_payload["frame"],
                list_values=last_payload["list_values"],
                file_name=export_name_widget.value,
                default_stem=last_payload["default_stem"],
                file_format=export_format_widget.value,
            )
            with export_output:
                print(f"Exportado en: {path}")
        except Exception as exc:
            with export_output:
                print(f"No se pudo exportar la tabla PCA: {exc}")

    scaler_widget.observe(_sync_components, names="value")
    refresh_button.on_click(_render)
    export_button.on_click(_export)
    _sync_components()
    _render()

    controls = widgets.VBox(
        [
            widgets.HTML("<b>Tabla interactiva de PCA</b>"),
            widgets.HTML("<i>Misma calidad analitica: solo cambias la vista, no el motor PCA.</i>"),
            scaler_widget,
            view_widget,
            component_widget,
            top_n_widget,
            score_rows_widget,
            sort_widget,
            widgets.HBox([ascending_widget, round_widget]),
            refresh_button,
            widgets.HTML("<b>Exportacion</b>"),
            export_name_widget,
            widgets.HBox([export_format_widget, export_button]),
            export_output,
            status_output,
        ]
    )
    content = widgets.VBox([summary_output, table_output], layout=widgets.Layout(width="100%"))
    return widgets.HBox([controls, content], layout=widgets.Layout(align_items="flex-start"))


def _clustering_interactive_table(
    self,
    top_n=10,
    compact_max_features=5,
    compact_criterion="holdout_silhouette",
    test_size=0.25,
    n_splits=2,
    min_cluster_share=0.01,
    include_stability=True,
    stability_sample_fraction=0.8,
    stability_n_splits=3,
):
    self._require_fit()
    _require_widgets()

    raw_frame = self._as_frame(self.data)
    raw_columns = [column for column in raw_frame.columns if column != self.target_name_]
    default_groupby = tuple(
        column
        for column in self.list_influential_variables(top_n=min(5, len(raw_columns)))
        if column in raw_columns
    )
    if not default_groupby:
        default_groupby = tuple(raw_columns[: min(5, len(raw_columns))])

    analysis_state = {"signature": None, "report": None}
    max_compact_features = max(2, len(self.list_influential_variables(top_n=None)))

    view_widget = widgets.Dropdown(
        options=[
            ("Mejor configuracion", "best_configuration"),
            ("Leaderboard", "leaderboard"),
            ("Comparativa por escalado", "scaler_report"),
            ("Metricas", "metric_report"),
            ("Tamano de clusters", "cluster_size_report"),
            ("Groupby por cluster", "cluster_groupby"),
            ("Cluster vs target", "cluster_target_report"),
            ("Dataframe clusterizado", "clustered_dataframe"),
            ("Holdout", "holdout_top"),
            ("Subconjuntos compactos", "compact_subset_search"),
            ("Groupby del mejor compacto", "best_compact_groupby"),
            ("Estabilidad", "stability_top"),
            ("Resumen de estabilidad", "stability_summary"),
        ],
        value="leaderboard",
        description="Vista",
        layout=widgets.Layout(width="280px"),
    )
    top_n_widget = widgets.IntSlider(
        value=min(max(3, int(top_n)), 20),
        min=3,
        max=20,
        step=1,
        description="Top n",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    rows_widget = widgets.IntSlider(
        value=10,
        min=5,
        max=30,
        step=1,
        description="Filas",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    groupby_variables_widget = widgets.SelectMultiple(
        options=raw_columns,
        value=default_groupby,
        description="Variables",
        rows=min(8, max(4, len(default_groupby))),
        layout=widgets.Layout(width="260px", height="190px"),
    )
    include_target_widget = widgets.Checkbox(value=True, description="Incluir target")
    numeric_agg_widget = widgets.Dropdown(
        options=[("Media", "mean"), ("Mediana", "median")],
        value="mean",
        description="Numerica",
        layout=widgets.Layout(width="220px"),
    )
    categorical_agg_widget = widgets.Dropdown(
        options=[("Moda", "mode"), ("Primero", "first")],
        value="mode",
        description="Categorica",
        layout=widgets.Layout(width="220px"),
    )
    normalize_widget = widgets.Dropdown(
        options=[
            ("Por cluster", "index"),
            ("Por target", "columns"),
            ("Global", "all"),
            ("Sin normalizar", False),
        ],
        value="index",
        description="Normalize",
        layout=widgets.Layout(width="220px"),
    )
    include_components_widget = widgets.Checkbox(value=True, description="Anadir PCs")
    component_widget = widgets.IntSlider(
        value=max(2, min(3, int(self.best_configuration()["n_components"]))),
        min=2,
        max=max(2, int(self.best_configuration()["n_components"])),
        step=1,
        description="PCs out",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    round_widget = widgets.Dropdown(
        options=[("Sin redondeo", None)] + [(str(value), value) for value in range(0, 7)],
        value=4,
        description="Round",
        layout=widgets.Layout(width="220px"),
    )
    compact_criterion_widget = widgets.Dropdown(
        options=[
            ("Holdout silhouette", "holdout_silhouette"),
            ("Holdout davies bouldin", "holdout_davies_bouldin"),
            ("Silhouette interno", "silhouette"),
            ("Davies bouldin interno", "davies_bouldin"),
        ],
        value=compact_criterion,
        description="Compacto",
        layout=widgets.Layout(width="280px"),
    )
    compact_max_features_widget = widgets.IntSlider(
        value=min(max(2, int(compact_max_features)), max_compact_features),
        min=2,
        max=max_compact_features,
        step=1,
        description="Max feats",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    test_size_widget = widgets.FloatSlider(
        value=float(test_size),
        min=0.1,
        max=0.5,
        step=0.05,
        readout_format=".2f",
        description="Test size",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    n_splits_widget = widgets.IntSlider(
        value=int(n_splits),
        min=1,
        max=6,
        step=1,
        description="Splits",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    min_cluster_share_widget = widgets.FloatSlider(
        value=float(min_cluster_share),
        min=0.0,
        max=0.1,
        step=0.005,
        readout_format=".3f",
        description="Min share",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    include_stability_widget = widgets.Checkbox(value=bool(include_stability), description="Con estabilidad")
    stability_sample_fraction_widget = widgets.FloatSlider(
        value=float(stability_sample_fraction),
        min=0.5,
        max=0.95,
        step=0.05,
        readout_format=".2f",
        description="Sample frac",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    stability_n_splits_widget = widgets.IntSlider(
        value=int(stability_n_splits),
        min=1,
        max=6,
        step=1,
        description="Stab splits",
        continuous_update=False,
        layout=widgets.Layout(width="280px"),
    )
    export_name_widget, export_format_widget, export_button = _build_export_controls(
        "fenrir_clustering.xlsx",
        "Exportar vista",
    )
    refresh_button = widgets.Button(
        description="Actualizar clustering",
        button_style="primary",
        icon="refresh",
    )

    status_output = widgets.Output()
    summary_output = widgets.Output()
    table_output = widgets.Output()
    export_output = widgets.Output()
    last_payload = {"frame": None, "list_values": None, "default_stem": "fenrir_clustering"}

    def _analysis_signature():
        return (
            int(top_n_widget.value),
            int(compact_max_features_widget.value),
            compact_criterion_widget.value,
            float(test_size_widget.value),
            int(n_splits_widget.value),
            float(min_cluster_share_widget.value),
            bool(include_stability_widget.value),
            float(stability_sample_fraction_widget.value),
            int(stability_n_splits_widget.value),
        )

    def _get_analysis_report():
        signature = _analysis_signature()
        if analysis_state["report"] is None or analysis_state["signature"] != signature:
            analysis_state["report"] = self.analysis_report(
                top_n=int(top_n_widget.value),
                leaderboard_top_n=int(top_n_widget.value),
                holdout_top_n=int(top_n_widget.value),
                compact_max_features=int(compact_max_features_widget.value),
                compact_criterion=compact_criterion_widget.value,
                test_size=float(test_size_widget.value),
                n_splits=int(n_splits_widget.value),
                stratify_target=True,
                min_cluster_share=float(min_cluster_share_widget.value),
                include_stability=bool(include_stability_widget.value),
                stability_sample_fraction=float(stability_sample_fraction_widget.value),
                stability_n_splits=int(stability_n_splits_widget.value),
                stability_top_n=int(top_n_widget.value),
                stability_stratify_target=True,
            )
            analysis_state["signature"] = signature
        return analysis_state["report"]

    def _render(_=None):
        with status_output:
            clear_output(wait=True)
            print("Actualizando tabla de clustering...")

        try:
            view = view_widget.value
            requires_analysis = view in {
                "holdout_top",
                "compact_subset_search",
                "best_compact_groupby",
                "stability_top",
                "stability_summary",
            }
            report = _get_analysis_report() if requires_analysis else None

            summary = pd.DataFrame(
                {
                    "valor": [
                        view,
                        self.best_configuration()["scaler_name"],
                        self.best_configuration()["algorithm"],
                        int(self.best_configuration()["n_clusters"]),
                        int(self.best_configuration()["n_components"]),
                    ]
                },
                index=[
                    "vista_activa",
                    "scaler_base",
                    "algoritmo_base",
                    "clusters_base",
                    "componentes_base",
                ],
            )

            list_values = None
            if view == "best_configuration":
                frame = self.best_configuration().rename("valor").to_frame()
            elif view == "leaderboard":
                frame = self.summary(top_n=int(top_n_widget.value))
            elif view == "scaler_report":
                frame = self.scaler_report()
            elif view == "metric_report":
                frame = self.metric_report()
            elif view == "cluster_size_report":
                frame = self.cluster_size_report()
            elif view == "cluster_groupby":
                selected_variables = list(groupby_variables_widget.value) or None
                frame = self.cluster_groupby(
                    variables=selected_variables,
                    top_n=int(top_n_widget.value),
                    include_target=bool(include_target_widget.value),
                    numeric_agg=numeric_agg_widget.value,
                    categorical_agg=categorical_agg_widget.value,
                    round_digits=int(round_widget.value if round_widget.value is not None else 4),
                )
            elif view == "cluster_target_report":
                frame = self.cluster_target_report(normalize=normalize_widget.value)
            elif view == "clustered_dataframe":
                frame = self.clustered_dataframe(
                    include_components=bool(include_components_widget.value),
                    n_components=int(component_widget.value),
                ).head(int(rows_widget.value))
            elif view == "holdout_top":
                frame = report["holdout_top"].head(int(top_n_widget.value))
            elif view == "compact_subset_search":
                frame = report["compact_subset_search"].head(int(top_n_widget.value))
                list_values = list(report.get("best_compact_variables", []))
            elif view == "best_compact_groupby":
                frame = report["best_compact_groupby"]
                list_values = list(report.get("best_compact_variables", []))
            elif view == "stability_top":
                frame = report["stability_top"].head(int(top_n_widget.value))
            else:
                frame = report["stability_summary"]

            frame = _round_frame(frame, round_widget.value)
            last_payload["frame"] = frame.copy()
            last_payload["list_values"] = list_values
            last_payload["default_stem"] = f"fenrir_clustering_{view}"

            with summary_output:
                clear_output(wait=True)
                display(summary)

            with table_output:
                clear_output(wait=True)
                display(frame)

            with status_output:
                clear_output(wait=True)
                print("Tabla de clustering lista")
        except Exception as exc:
            last_payload["frame"] = None
            last_payload["list_values"] = None
            with summary_output:
                clear_output(wait=True)
            with table_output:
                clear_output(wait=True)
            with status_output:
                clear_output(wait=True)
                print(f"No se pudo generar la tabla de clustering: {exc}")

    def _export(_=None):
        with export_output:
            clear_output(wait=True)
        try:
            if last_payload["frame"] is None:
                raise RuntimeError("Actualiza la tabla de clustering antes de exportar.")
            path = _export_payload(
                frame=last_payload["frame"],
                list_values=last_payload["list_values"],
                file_name=export_name_widget.value,
                default_stem=last_payload["default_stem"],
                file_format=export_format_widget.value,
            )
            with export_output:
                print(f"Exportado en: {path}")
        except Exception as exc:
            with export_output:
                print(f"No se pudo exportar la vista de clustering: {exc}")

    refresh_button.on_click(_render)
    export_button.on_click(_export)
    _render()

    controls = widgets.VBox(
        [
            widgets.HTML("<b>Tabla interactiva de clustering</b>"),
            widgets.HTML("<i>Misma calidad analitica: la UI usa el mismo best fit, holdout, compactos y estabilidad de Fenrir.</i>"),
            view_widget,
            top_n_widget,
            rows_widget,
            groupby_variables_widget,
            widgets.HBox([include_target_widget, include_components_widget]),
            widgets.HBox([numeric_agg_widget, categorical_agg_widget]),
            widgets.HBox([normalize_widget, component_widget]),
            compact_criterion_widget,
            compact_max_features_widget,
            test_size_widget,
            n_splits_widget,
            min_cluster_share_widget,
            widgets.HBox([include_stability_widget, stability_n_splits_widget]),
            stability_sample_fraction_widget,
            round_widget,
            refresh_button,
            widgets.HTML("<b>Exportacion</b>"),
            export_name_widget,
            widgets.HBox([export_format_widget, export_button]),
            export_output,
            status_output,
        ]
    )
    content = widgets.VBox([summary_output, table_output], layout=widgets.Layout(width="100%"))
    return widgets.HBox([controls, content], layout=widgets.Layout(align_items="flex-start"))


def _interactive_tables(
    self,
    pca_top_n=10,
    cluster_top_n=10,
    compact_max_features=5,
    compact_criterion="holdout_silhouette",
    test_size=0.25,
    n_splits=2,
    min_cluster_share=0.01,
    include_stability=True,
    stability_sample_fraction=0.8,
    stability_n_splits=3,
):
    _require_widgets()
    pca_panel = self.pca_interactive_table(top_n=pca_top_n)
    cluster_panel = self.clustering_interactive_table(
        top_n=cluster_top_n,
        compact_max_features=compact_max_features,
        compact_criterion=compact_criterion,
        test_size=test_size,
        n_splits=n_splits,
        min_cluster_share=min_cluster_share,
        include_stability=include_stability,
        stability_sample_fraction=stability_sample_fraction,
        stability_n_splits=stability_n_splits,
    )
    tabs = widgets.Tab(children=[pca_panel, cluster_panel])
    tabs.set_title(0, "PCA")
    tabs.set_title(1, "Clustering")
    return tabs


def _interactive_pca(self, **kwargs):
    return self.pca_interactive_table(**kwargs)


def _interactive_clustering(self, **kwargs):
    return self.clustering_interactive_table(**kwargs)


def _interactive_workbench(self, data=None, fenrir_config=None, analysis_config=None):
    self._require_fit()
    _require_widgets()

    raw_data = self._as_frame(self.data if data is None else data)
    available_columns = list(raw_data.columns)
    available_scalers = _default_scaler_catalog()
    default_preset_file = "fenrir_workbench_presets.json"

    initial_fenrir_config, initial_analysis_input = _prepare_workbench_initial_configs(
        self,
        data=data,
        fenrir_config=fenrir_config,
        analysis_config=analysis_config,
    )
    initial_analysis_config = _default_analysis_config(self, initial_analysis_input)

    preset_store = copy.deepcopy(_workbench_presets_for(self))
    pending_file_message = None
    try:
        persisted_store, persisted_path = _load_preset_store_json(
            file_name=default_preset_file,
            default_stem="fenrir_workbench_presets",
        )
        if persisted_store:
            preset_store.update(persisted_store)
            pending_file_message = f"Presets JSON cargados automaticamente desde: {persisted_path}"
    except FileNotFoundError:
        pass

    self.workbench_presets_ = copy.deepcopy(preset_store)
    state = {
        "model": self,
        "report": None,
        "current_figure": None,
        "preset_store": copy.deepcopy(preset_store),
    }

    target_widget = widgets.Dropdown(
        options=[("Sin target", "__none__")] + [(column, column) for column in available_columns],
        value="__none__" if initial_fenrir_config["target"] is None else initial_fenrir_config["target"],
        description="Target",
        layout=widgets.Layout(width="280px"),
    )
    use_features_widget = widgets.Checkbox(value=bool(initial_fenrir_config["features"]), description="Usar features")
    features_widget = widgets.SelectMultiple(
        options=available_columns,
        value=tuple(initial_fenrir_config["features"]),
        description="Features",
        rows=min(10, max(5, len(available_columns))),
        layout=widgets.Layout(width="300px", height="220px"),
    )
    exclude_widget = widgets.SelectMultiple(
        options=available_columns,
        value=tuple(initial_fenrir_config["exclude_columns"]),
        description="Excluir",
        rows=min(10, max(5, len(available_columns))),
        layout=widgets.Layout(width="300px", height="220px"),
    )
    variance_widget = widgets.FloatSlider(
        value=float(initial_fenrir_config["variance_threshold"]),
        min=0.6,
        max=0.99,
        step=0.01,
        readout_format=".2f",
        description="Varianza",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    cluster_min_widget = widgets.IntSlider(
        value=int(min(initial_fenrir_config["cluster_range"])),
        min=2,
        max=12,
        step=1,
        description="Cluster min",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    cluster_max_widget = widgets.IntSlider(
        value=int(max(initial_fenrir_config["cluster_range"])),
        min=2,
        max=15,
        step=1,
        description="Cluster max",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    algorithms_widget = widgets.SelectMultiple(
        options=[("KMeans", "kmeans"), ("GMM", "gmm"), ("Agglomerative", "agglomerative")],
        value=tuple(initial_fenrir_config["algorithms"]),
        description="Algoritmos",
        rows=3,
        layout=widgets.Layout(width="280px", height="125px"),
    )
    scalers_widget = widgets.SelectMultiple(
        options=list(available_scalers.keys()),
        value=tuple(initial_fenrir_config["scalers"]),
        description="Scalers",
        rows=3,
        layout=widgets.Layout(width="280px", height="125px"),
    )
    max_categories_widget = widgets.IntSlider(
        value=int(initial_fenrir_config["max_categories"]),
        min=2,
        max=50,
        step=1,
        description="Max cat",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    random_state_widget = widgets.IntText(
        value=int(initial_fenrir_config["random_state"]),
        description="Seed",
        layout=widgets.Layout(width="220px"),
    )

    top_n_widget = widgets.IntSlider(
        value=int(initial_analysis_config["top_n"]),
        min=3,
        max=20,
        step=1,
        description="Top n",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    leaderboard_top_n_widget = widgets.IntSlider(
        value=int(initial_analysis_config["leaderboard_top_n"]),
        min=3,
        max=25,
        step=1,
        description="Leaderboard",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    compact_max_features_widget = widgets.IntSlider(
        value=int(initial_analysis_config["compact_max_features"]),
        min=2,
        max=max(2, len(available_columns)),
        step=1,
        description="Compact max",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    compact_criterion_widget = widgets.Dropdown(
        options=[
            ("Holdout silhouette", "holdout_silhouette"),
            ("Holdout davies bouldin", "holdout_davies_bouldin"),
            ("Silhouette interno", "silhouette"),
            ("Davies bouldin interno", "davies_bouldin"),
        ],
        value=initial_analysis_config["compact_criterion"],
        description="Compacto",
        layout=widgets.Layout(width="320px"),
    )
    test_size_widget = widgets.FloatSlider(
        value=float(initial_analysis_config["test_size"]),
        min=0.1,
        max=0.5,
        step=0.05,
        readout_format=".2f",
        description="Test size",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    n_splits_widget = widgets.IntSlider(
        value=int(initial_analysis_config["n_splits"]),
        min=1,
        max=6,
        step=1,
        description="Splits",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    min_cluster_share_widget = widgets.FloatSlider(
        value=float(initial_analysis_config["min_cluster_share"]),
        min=0.0,
        max=0.1,
        step=0.005,
        readout_format=".3f",
        description="Min share",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    include_stability_widget = widgets.Checkbox(value=bool(initial_analysis_config["include_stability"]), description="Con estabilidad")
    stability_sample_fraction_widget = widgets.FloatSlider(
        value=float(initial_analysis_config["stability_sample_fraction"]),
        min=0.5,
        max=0.95,
        step=0.05,
        readout_format=".2f",
        description="Sample frac",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    stability_n_splits_widget = widgets.IntSlider(
        value=int(initial_analysis_config["stability_n_splits"]),
        min=1,
        max=6,
        step=1,
        description="Stab splits",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )

    plot_type_widget = widgets.Dropdown(
        options=[
            ("Scree", "scree"),
            ("Circulo de correlaciones", "correlation_circle"),
            ("Contribuciones", "feature_contributions"),
            ("Calidad de variables", "feature_quality"),
            ("Clusters", "clusters"),
        ],
        value="scree",
        description="Grafico",
        layout=widgets.Layout(width="320px"),
    )
    plot_scaler_widget = widgets.Dropdown(
        options=[(name, name) for name in state["model"].scalers.keys()],
        value=state["model"].best_configuration()["scaler_name"],
        description="Scaler plot",
        layout=widgets.Layout(width="320px"),
    )
    plot_top_n_widget = widgets.IntSlider(
        value=max(5, int(initial_analysis_config["top_n"])),
        min=3,
        max=25,
        step=1,
        description="Plot top n",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    plot_width_widget = widgets.FloatSlider(
        value=8.0,
        min=5.0,
        max=16.0,
        step=0.5,
        readout_format=".1f",
        description="Width",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    plot_height_widget = widgets.FloatSlider(
        value=5.5,
        min=4.0,
        max=12.0,
        step=0.5,
        readout_format=".1f",
        description="Height",
        continuous_update=False,
        layout=widgets.Layout(width="320px"),
    )
    plot_export_name_widget = widgets.Text(
        value="fenrir_plot.png",
        description="Archivo",
        layout=widgets.Layout(width="320px"),
    )
    plot_export_format_widget = widgets.Dropdown(
        options=[("PNG", "png"), ("SVG", "svg"), ("PDF", "pdf")],
        value="png",
        description="Formato",
        layout=widgets.Layout(width="220px"),
    )
    plot_refresh_button = widgets.Button(description="Actualizar grafico", button_style="primary", icon="image")
    plot_export_button = widgets.Button(description="Exportar grafico", icon="download")

    preset_name_widget = widgets.Text(
        value="",
        placeholder="mi_configuracion",
        description="Preset",
        layout=widgets.Layout(width="320px"),
    )
    preset_dropdown = widgets.Dropdown(
        options=[("Sin preset", "__none__")],
        value="__none__",
        description="Guardados",
        layout=widgets.Layout(width="320px"),
    )
    preset_save_button = widgets.Button(description="Guardar", icon="save")
    preset_load_button = widgets.Button(description="Cargar", icon="upload")
    preset_delete_button = widgets.Button(description="Borrar", icon="trash")

    view_mode_widget = widgets.Dropdown(
        options=[
            ("PCA + clustering", "tabs"),
            ("Solo PCA", "pca"),
            ("Solo clustering", "clustering"),
            ("Comparar presets", "compare"),
        ],
        value="tabs",
        description="Vista",
        layout=widgets.Layout(width="320px"),
    )
    compare_left_widget = widgets.Dropdown(
        options=[("Actual", "__current__")],
        value="__current__",
        description="Izquierda",
        layout=widgets.Layout(width="320px"),
    )
    compare_right_widget = widgets.Dropdown(
        options=[("Actual", "__current__")],
        value="__current__",
        description="Derecha",
        layout=widgets.Layout(width="320px"),
    )
    compare_button = widgets.Button(description="Comparar", button_style="primary", icon="columns")

    persistence_file_widget = widgets.Text(
        value=default_preset_file,
        description="Presets JSON",
        layout=widgets.Layout(width="330px"),
    )
    auto_sync_widget = widgets.Checkbox(value=True, description="Auto-sync")
    merge_load_widget = widgets.Checkbox(value=True, description="Combinar al cargar")
    save_json_button = widgets.Button(description="Guardar JSON", icon="save")
    load_json_button = widgets.Button(description="Cargar JSON", icon="folder-open")
    executive_file_widget = widgets.Text(
        value="fenrir_executivo.xlsx",
        description="Excel",
        layout=widgets.Layout(width="330px"),
    )
    export_exec_button = widgets.Button(description="Exportar ejecutivo", icon="download")

    refresh_button = widgets.Button(description="Reajustar Fenrir", button_style="primary", icon="refresh")
    deep_validate_button = widgets.Button(description="Validar profundo", button_style="warning", icon="check")

    status_output = widgets.Output()
    quality_output = widgets.Output()
    validation_output = widgets.Output()
    tables_output = widgets.Output()
    plot_output = widgets.Output()
    plot_status_output = widgets.Output()
    preset_output = widgets.Output()
    compare_status_output = widgets.Output()
    file_output = widgets.Output()

    def _write_output(output, message):
        with output:
            clear_output(wait=True)
            print(message)

    def _sync_preset_store(model):
        model.workbench_presets_ = copy.deepcopy(state["preset_store"])
        _publish_notebook_bindings(fenrir_workbench_presets=copy.deepcopy(state["preset_store"]))

    def _refresh_preset_options(selected=None):
        options = [("Sin preset", "__none__")] + [(name, name) for name in sorted(state["preset_store"])]
        preset_dropdown.options = options
        option_values = {value for _, value in options}
        if selected in option_values:
            preset_dropdown.value = selected
        elif preset_dropdown.value not in option_values:
            preset_dropdown.value = "__none__"

    def _refresh_compare_options(selected_left=None, selected_right=None):
        options = [("Actual", "__current__")] + [(name, name) for name in sorted(state["preset_store"])]
        compare_left_widget.options = options
        compare_right_widget.options = options
        option_values = {value for _, value in options}

        if selected_left in option_values:
            compare_left_widget.value = selected_left
        elif compare_left_widget.value not in option_values:
            compare_left_widget.value = "__current__"

        if selected_right in option_values and selected_right != compare_left_widget.value:
            compare_right_widget.value = selected_right
        elif compare_right_widget.value not in option_values or compare_right_widget.value == compare_left_widget.value:
            compare_right_widget.value = next(
                (value for _, value in options if value != compare_left_widget.value),
                compare_left_widget.value,
            )

    def _sync_column_options(*_):
        selected_target = None if target_widget.value == "__none__" else target_widget.value
        available_feature_columns = [column for column in available_columns if column != selected_target]
        current_features = tuple(column for column in features_widget.value if column in available_feature_columns)
        current_excludes = tuple(column for column in exclude_widget.value if column in available_feature_columns)
        features_widget.options = available_feature_columns
        exclude_widget.options = available_feature_columns
        features_widget.value = current_features
        exclude_widget.value = current_excludes
        compact_max_features_widget.max = max(2, len(available_feature_columns) or 2)
        compact_max_features_widget.value = min(compact_max_features_widget.value, compact_max_features_widget.max)

    def _current_fenrir_config():
        selected_target = None if target_widget.value == "__none__" else target_widget.value
        selected_features = list(features_widget.value) if use_features_widget.value else None
        selected_excludes = list(exclude_widget.value)
        selected_algorithms = tuple(algorithms_widget.value)
        selected_scalers = tuple(scalers_widget.value)
        cluster_min = int(cluster_min_widget.value)
        cluster_max = int(cluster_max_widget.value)

        if cluster_max < cluster_min:
            raise ValueError("cluster_max no puede ser menor que cluster_min.")
        if not selected_algorithms:
            raise ValueError("Debes seleccionar al menos un algoritmo.")
        if not selected_scalers:
            raise ValueError("Debes seleccionar al menos un scaler.")

        return {
            "target": selected_target,
            "features": selected_features,
            "exclude_columns": selected_excludes,
            "variance_threshold": float(variance_widget.value),
            "cluster_range": range(cluster_min, cluster_max + 1),
            "algorithms": selected_algorithms,
            "max_categories": int(max_categories_widget.value),
            "random_state": int(random_state_widget.value),
            "scalers": {name: clone(available_scalers[name]) for name in selected_scalers},
        }

    def _current_analysis_config():
        top_n_value = int(top_n_widget.value)
        return {
            "top_n": top_n_value,
            "leaderboard_top_n": int(leaderboard_top_n_widget.value),
            "holdout_top_n": top_n_value,
            "compact_max_features": int(compact_max_features_widget.value),
            "compact_criterion": compact_criterion_widget.value,
            "compact_include_holdout": None,
            "test_size": float(test_size_widget.value),
            "n_splits": int(n_splits_widget.value),
            "stratify_target": True,
            "min_cluster_share": float(min_cluster_share_widget.value),
            "include_stability": bool(include_stability_widget.value),
            "stability_sample_fraction": float(stability_sample_fraction_widget.value),
            "stability_n_splits": int(stability_n_splits_widget.value),
            "stability_top_n": top_n_value,
            "stability_stratify_target": True,
        }

    def _normalized_analysis():
        return _default_analysis_config(state["model"], _current_analysis_config())

    def _sync_plot_scalers(model):
        options = [(name, name) for name in model.scalers.keys()]
        best_scaler = model.best_configuration()["scaler_name"]
        plot_scaler_widget.options = options
        plot_scaler_widget.value = best_scaler if best_scaler in dict(options) else options[0][1]

    def _render_plot(_=None):
        _write_output(plot_status_output, "Actualizando grafico...")
        try:
            if state["current_figure"] is not None:
                    _require_matplotlib_pyplot().close(state["current_figure"])
            figure = _render_plot_figure(
                state["model"],
                plot_type_widget.value,
                scaler_name=plot_scaler_widget.value,
                top_n=int(plot_top_n_widget.value),
                width=float(plot_width_widget.value),
                height=float(plot_height_widget.value),
            )
            state["current_figure"] = figure
            with plot_output:
                clear_output(wait=True)
                display(figure)
            _write_output(plot_status_output, "Grafico listo")
        except Exception as exc:
            with plot_output:
                clear_output(wait=True)
            _write_output(plot_status_output, f"No se pudo generar el grafico: {exc}")

    def _export_plot(_=None):
        try:
            if state["current_figure"] is None:
                raise RuntimeError("Actualiza el grafico antes de exportar.")
            path = _export_figure(
                state["current_figure"],
                file_name=plot_export_name_widget.value,
                default_stem=f"fenrir_plot_{plot_type_widget.value}",
                file_format=plot_export_format_widget.value,
            )
            _write_output(plot_status_output, f"Grafico exportado en: {path}")
        except Exception as exc:
            _write_output(plot_status_output, f"No se pudo exportar el grafico: {exc}")

    def _render_selected_view(_=None):
        current_model = state["model"]
        current_fenrir = _current_fenrir_config()
        current_analysis = _normalized_analysis()
        view_mode = view_mode_widget.value

        with compare_status_output:
            clear_output(wait=True)

        if view_mode == "compare":
            _write_output(compare_status_output, "Preparando comparacion...")
            try:
                comparison = current_model.compare_workbench_candidates(
                    left_candidate=compare_left_widget.value,
                    right_candidate=compare_right_widget.value,
                    preset_store=copy.deepcopy(state["preset_store"]),
                    current_fenrir_config=current_fenrir,
                    current_analysis_config=current_analysis,
                    data=data,
                )
                with tables_output:
                    clear_output(wait=True)
                    print(f"Comparacion: {comparison['left_label']} vs {comparison['right_label']}")
                    print("Configuracion")
                    display(comparison["configuration"])
                    print("Mejor ajuste")
                    display(comparison["best_configuration"])
                    print("Calidad")
                    display(comparison["quality"])
                    print("Validacion")
                    display(comparison["validation"])
                _publish_notebook_bindings(fenrir_comparison=comparison)
                _write_output(compare_status_output, "Comparacion lista")
            except Exception as exc:
                with tables_output:
                    clear_output(wait=True)
                _write_output(compare_status_output, f"No se pudo comparar: {exc}")
            return

        if view_mode == "tabs":
            panel = current_model.interactive_tables(
                pca_top_n=max(10, int(current_analysis["top_n"])),
                cluster_top_n=max(10, int(current_analysis["leaderboard_top_n"])),
                compact_max_features=int(current_analysis["compact_max_features"]),
                compact_criterion=current_analysis["compact_criterion"],
                test_size=float(current_analysis["test_size"]),
                n_splits=int(current_analysis["n_splits"]),
                min_cluster_share=float(current_analysis["min_cluster_share"]),
                include_stability=bool(current_analysis["include_stability"]),
                stability_sample_fraction=float(current_analysis["stability_sample_fraction"]),
                stability_n_splits=int(current_analysis["stability_n_splits"]),
            )
        elif view_mode == "pca":
            panel = current_model.pca_interactive_table(top_n=max(10, int(current_analysis["top_n"])))
        else:
            panel = current_model.clustering_interactive_table(
                top_n=max(10, int(current_analysis["leaderboard_top_n"])),
                compact_max_features=int(current_analysis["compact_max_features"]),
                compact_criterion=current_analysis["compact_criterion"],
                test_size=float(current_analysis["test_size"]),
                n_splits=int(current_analysis["n_splits"]),
                min_cluster_share=float(current_analysis["min_cluster_share"]),
                include_stability=bool(current_analysis["include_stability"]),
                stability_sample_fraction=float(current_analysis["stability_sample_fraction"]),
                stability_n_splits=int(current_analysis["stability_n_splits"]),
            )

        with tables_output:
            clear_output(wait=True)
            display(panel)
        _publish_notebook_bindings(fenrir_tables=panel)

    def _render_workbench(_=None):
        _write_output(status_output, "Reajustando Fenrir con la misma logica analitica...")
        try:
            current_fenrir = _current_fenrir_config()
            current_analysis = _current_analysis_config()
            model = Fenrir(raw_data.copy(), **current_fenrir)
            model.fit()
            state["model"] = model
            state["report"] = None
            _sync_preset_store(model)

            with quality_output:
                clear_output(wait=True)
                print("Control de calidad del ajuste actual")
                display(model.best_configuration().rename("valor").to_frame())
                display(_key_metric_view(model))

            with validation_output:
                clear_output(wait=True)
                print(
                    "Validacion profunda pendiente. Usa el boton para forzar holdout, compactos y estabilidad con la configuracion actual."
                )

            _sync_plot_scalers(model)
            _refresh_compare_options(compare_left_widget.value, compare_right_widget.value)
            _render_plot()
            _render_selected_view()
            _publish_notebook_bindings(
                fenrir=model,
                clusters=model.best_labels_.copy(),
                fenrir_workbench_model=model,
                fenrir_config=copy.deepcopy(current_fenrir),
                analysis_config=copy.deepcopy(current_analysis),
                report=None,
            )
            _write_output(
                status_output,
                "Fenrir reajustado. La calidad se conserva porque el workbench usa el mismo pipeline de ajuste y validacion.",
            )
        except Exception as exc:
            with quality_output:
                clear_output(wait=True)
            with validation_output:
                clear_output(wait=True)
            with tables_output:
                clear_output(wait=True)
            with plot_output:
                clear_output(wait=True)
            _write_output(status_output, f"No se pudo reajustar Fenrir: {exc}")

    def _run_deep_validation(_=None):
        _write_output(status_output, "Ejecutando validacion profunda...")
        try:
            current_analysis = _normalized_analysis()
            report = state["model"].analysis_report(**current_analysis)
            state["report"] = report
            holdout_view = report["holdout_top"].loc[
                :,
                [
                    column
                    for column in [
                        "scaler_name",
                        "algorithm",
                        "n_clusters",
                        "n_components",
                        "test_silhouette_mean",
                        "test_davies_bouldin_mean",
                        "test_v_measure_mean",
                    ]
                    if column in report["holdout_top"].columns
                ],
            ]
            compact_view = report["compact_subset_search"].loc[
                :,
                [
                    column
                    for column in [
                        "n_selected_features",
                        "selected_variables",
                        "test_silhouette_mean",
                        "test_davies_bouldin_mean",
                        "silhouette",
                        "min_cluster_size",
                        "selection_status",
                    ]
                    if column in report["compact_subset_search"].columns
                ],
            ]

            with validation_output:
                clear_output(wait=True)
                print("Validacion profunda completada")
                print("Mejor lista compacta:", report["best_compact_variables"])
                print("Top holdout")
                display(holdout_view.head(int(top_n_widget.value)))
                print("Compactos evaluados")
                display(compact_view.head(int(top_n_widget.value)))
                if report.get("stability_summary") is not None:
                    print("Resumen de estabilidad")
                    display(report["stability_summary"])

            _publish_notebook_bindings(report=report, analysis_config=copy.deepcopy(current_analysis))
            _render_selected_view()
            _write_output(status_output, "Validacion profunda completada con holdout, compactos y estabilidad.")
        except Exception as exc:
            with validation_output:
                clear_output(wait=True)
            _write_output(status_output, f"No se pudo ejecutar la validacion profunda: {exc}")

    def _apply_preset_to_widgets(preset):
        fenrir_payload = copy.deepcopy(preset.get("fenrir_config", {}))
        analysis_payload = copy.deepcopy(preset.get("analysis_config", {}))

        target_widget.value = "__none__" if fenrir_payload.get("target") is None else fenrir_payload.get("target")
        _sync_column_options()
        use_features_widget.value = fenrir_payload.get("features") is not None
        valid_features = tuple(column for column in fenrir_payload.get("features") or [] if column in features_widget.options)
        valid_excludes = tuple(column for column in fenrir_payload.get("exclude_columns") or [] if column in exclude_widget.options)
        features_widget.value = valid_features
        exclude_widget.value = valid_excludes
        variance_widget.value = float(fenrir_payload.get("variance_threshold", variance_widget.value))

        preset_cluster_range = tuple(fenrir_payload.get("cluster_range") or range(cluster_min_widget.value, cluster_max_widget.value + 1))
        cluster_min_widget.value = int(min(preset_cluster_range))
        cluster_max_widget.value = int(max(preset_cluster_range))

        valid_algorithms = tuple(value for value in fenrir_payload.get("algorithms") or () if value in dict(algorithms_widget.options).values())
        valid_scalers = tuple(value for value in fenrir_payload.get("scalers") or () if value in scalers_widget.options)
        if valid_algorithms:
            algorithms_widget.value = valid_algorithms
        if valid_scalers:
            scalers_widget.value = valid_scalers

        max_categories_widget.value = int(fenrir_payload.get("max_categories", max_categories_widget.value))
        random_state_widget.value = int(fenrir_payload.get("random_state", random_state_widget.value))
        top_n_widget.value = int(analysis_payload.get("top_n", top_n_widget.value))
        leaderboard_top_n_widget.value = int(analysis_payload.get("leaderboard_top_n", leaderboard_top_n_widget.value))
        compact_max_features_widget.value = min(
            compact_max_features_widget.max,
            int(analysis_payload.get("compact_max_features", compact_max_features_widget.value)),
        )
        compact_criterion_widget.value = analysis_payload.get("compact_criterion", compact_criterion_widget.value)
        test_size_widget.value = float(analysis_payload.get("test_size", test_size_widget.value))
        n_splits_widget.value = int(analysis_payload.get("n_splits", n_splits_widget.value))
        min_cluster_share_widget.value = float(analysis_payload.get("min_cluster_share", min_cluster_share_widget.value))
        include_stability_widget.value = bool(analysis_payload.get("include_stability", include_stability_widget.value))
        stability_sample_fraction_widget.value = float(
            analysis_payload.get("stability_sample_fraction", stability_sample_fraction_widget.value)
        )
        stability_n_splits_widget.value = int(analysis_payload.get("stability_n_splits", stability_n_splits_widget.value))

    def _save_preset(_=None):
        try:
            preset_name = str(preset_name_widget.value or "").strip()
            if not preset_name:
                raise ValueError("Escribe un nombre para guardar el preset.")

            state["preset_store"][preset_name] = {
                "fenrir_config": copy.deepcopy(_current_fenrir_config()),
                "analysis_config": copy.deepcopy(_current_analysis_config()),
            }
            _sync_preset_store(state["model"])
            _refresh_preset_options(selected=preset_name)
            _refresh_compare_options(compare_left_widget.value, compare_right_widget.value)
            if auto_sync_widget.value:
                _save_json()
            _write_output(preset_output, f"Preset guardado: {preset_name}")
        except Exception as exc:
            _write_output(preset_output, f"No se pudo guardar el preset: {exc}")

    def _load_preset(_=None):
        try:
            if preset_dropdown.value == "__none__":
                raise ValueError("Selecciona un preset para cargar.")
            preset = state["preset_store"].get(preset_dropdown.value)
            if preset is None:
                raise ValueError("El preset seleccionado no existe.")

            _apply_preset_to_widgets(preset)
            preset_name_widget.value = preset_dropdown.value
            _write_output(preset_output, f"Preset cargado: {preset_dropdown.value}")
            _render_workbench()
        except Exception as exc:
            _write_output(preset_output, f"No se pudo cargar el preset: {exc}")

    def _delete_preset(_=None):
        try:
            if preset_dropdown.value == "__none__":
                raise ValueError("Selecciona un preset para borrar.")
            removed_name = preset_dropdown.value
            removed = state["preset_store"].pop(removed_name, None)
            if removed is None:
                raise ValueError("El preset seleccionado no existe.")

            _sync_preset_store(state["model"])
            _refresh_preset_options()
            _refresh_compare_options(compare_left_widget.value, compare_right_widget.value)
            if auto_sync_widget.value:
                _save_json()
            _render_selected_view()
            _write_output(preset_output, f"Preset borrado: {removed_name}")
        except Exception as exc:
            _write_output(preset_output, f"No se pudo borrar el preset: {exc}")

    def _save_json(_=None):
        try:
            path = _save_preset_store_json(
                state["preset_store"],
                file_name=persistence_file_widget.value,
                default_stem="fenrir_workbench_presets",
            )
            _write_output(file_output, f"Presets persistidos en: {path}")
        except Exception as exc:
            _write_output(file_output, f"No se pudieron persistir los presets: {exc}")

    def _load_json(_=None):
        try:
            loaded_store, path = _load_preset_store_json(
                file_name=persistence_file_widget.value,
                default_stem="fenrir_workbench_presets",
            )
            if merge_load_widget.value:
                merged_store = copy.deepcopy(state["preset_store"])
                merged_store.update(loaded_store)
                loaded_store = merged_store

            state["preset_store"] = copy.deepcopy(loaded_store)
            _sync_preset_store(state["model"])
            _refresh_preset_options()
            _refresh_compare_options(compare_left_widget.value, compare_right_widget.value)
            _render_selected_view()
            _write_output(file_output, f"Presets cargados desde: {path}")
        except Exception as exc:
            _write_output(file_output, f"No se pudieron cargar los presets: {exc}")

    def _export_executive(_=None):
        try:
            workbook_path = state["model"].export_executive_workbook(
                file_name=executive_file_widget.value,
                analysis_config=_current_analysis_config(),
                fenrir_config=_current_fenrir_config(),
            )
            _write_output(file_output, f"Excel ejecutivo exportado en: {workbook_path}")
        except Exception as exc:
            _write_output(file_output, f"No se pudo exportar el Excel ejecutivo: {exc}")

    target_widget.observe(_sync_column_options, names="value")
    refresh_button.on_click(_render_workbench)
    deep_validate_button.on_click(_run_deep_validation)
    plot_refresh_button.on_click(_render_plot)
    plot_export_button.on_click(_export_plot)
    preset_save_button.on_click(_save_preset)
    preset_load_button.on_click(_load_preset)
    preset_delete_button.on_click(_delete_preset)
    view_mode_widget.observe(_render_selected_view, names="value")
    compare_button.on_click(_render_selected_view)
    save_json_button.on_click(_save_json)
    load_json_button.on_click(_load_json)
    export_exec_button.on_click(_export_executive)

    model_box = widgets.VBox(
        [
            target_widget,
            use_features_widget,
            widgets.HBox([features_widget, exclude_widget]),
            variance_widget,
            widgets.HBox([cluster_min_widget, cluster_max_widget]),
            widgets.HBox([algorithms_widget, scalers_widget]),
            max_categories_widget,
            random_state_widget,
        ]
    )
    analysis_box = widgets.VBox(
        [
            top_n_widget,
            leaderboard_top_n_widget,
            compact_max_features_widget,
            compact_criterion_widget,
            test_size_widget,
            n_splits_widget,
            min_cluster_share_widget,
            include_stability_widget,
            stability_sample_fraction_widget,
            stability_n_splits_widget,
        ]
    )
    preset_box = widgets.VBox(
        [
            preset_name_widget,
            preset_dropdown,
            widgets.HBox([preset_save_button, preset_load_button, preset_delete_button]),
            preset_output,
        ]
    )
    plot_box = widgets.VBox(
        [
            plot_type_widget,
            plot_scaler_widget,
            plot_top_n_widget,
            plot_width_widget,
            plot_height_widget,
            plot_export_name_widget,
            widgets.HBox([plot_export_format_widget, plot_refresh_button, plot_export_button]),
            plot_status_output,
        ]
    )
    view_box = widgets.VBox(
        [
            widgets.HTML("<b>Vista de resultados</b>"),
            widgets.HTML("<i>No hace falta abrir las dos tablas a la vez: puedes ver solo PCA, solo clustering o comparar presets.</i>"),
            view_mode_widget,
        ]
    )
    compare_box = widgets.VBox(
        [
            widgets.HTML("<b>Comparacion lado a lado</b>"),
            widgets.HTML("<i>Compara el estado actual con uno o dos presets guardados usando la misma logica de ajuste y validacion.</i>"),
            compare_left_widget,
            compare_right_widget,
            compare_button,
            compare_status_output,
        ]
    )
    archive_box = widgets.VBox(
        [
            widgets.HTML("<b>Persistencia y exportacion ejecutiva</b>"),
            widgets.HTML("<i>Los presets pueden sobrevivir al reinicio del kernel en JSON y el informe ejecutivo genera un unico Excel con multiples hojas.</i>"),
            persistence_file_widget,
            widgets.HBox([auto_sync_widget, merge_load_widget]),
            widgets.HBox([save_json_button, load_json_button]),
            executive_file_widget,
            export_exec_button,
            file_output,
        ]
    )
    accordion = widgets.Accordion(children=[model_box, analysis_box, preset_box, plot_box, view_box, compare_box, archive_box])
    accordion.set_title(0, "Modelo")
    accordion.set_title(1, "Validacion")
    accordion.set_title(2, "Presets")
    accordion.set_title(3, "Graficos")
    accordion.set_title(4, "Vista")
    accordion.set_title(5, "Comparar")
    accordion.set_title(6, "Archivos")

    controls = widgets.VBox(
        [
            widgets.HTML("<b>Workbench interactivo de Fenrir</b>"),
            widgets.HTML("<i>La variable objetivo, las exportaciones, los presets y la comparacion usan exactamente el mismo pipeline validado del paquete.</i>"),
            accordion,
            widgets.HBox([refresh_button, deep_validate_button]),
            status_output,
        ],
        layout=widgets.Layout(width="440px"),
    )
    content = widgets.VBox([quality_output, validation_output, plot_output, tables_output], layout=widgets.Layout(width="100%"))

    _refresh_preset_options()
    _refresh_compare_options(selected_left="__current__")
    _sync_column_options()
    if pending_file_message:
        _write_output(file_output, pending_file_message)
    _render_workbench()
    return widgets.HBox([controls, content], layout=widgets.Layout(align_items="flex-start"))


def attach_interactive_api(cls=Fenrir):
    if getattr(cls, "_interactive_api_attached", False):
        return cls

    cls.pca_interactive_table = _pca_interactive_table
    cls.clustering_interactive_table = _clustering_interactive_table
    cls.interactive_tables = _interactive_tables
    cls.interactive_pca = _interactive_pca
    cls.interactive_clustering = _interactive_clustering
    cls.interactive_workbench = _interactive_workbench
    cls.save_workbench_preset = _save_workbench_preset
    cls.get_workbench_preset = _get_workbench_preset
    cls.delete_workbench_preset = _delete_workbench_preset
    cls.list_workbench_presets = _list_workbench_presets
    cls.save_workbench_presets_json = _save_workbench_presets_json
    cls.load_workbench_presets_json = _load_workbench_presets_json
    cls.export_executive_workbook = _export_executive_workbook
    cls.compare_workbench_candidates = _compare_workbench_candidates
    cls._interactive_api_attached = True
    return cls


attach_interactive_api(Fenrir)
